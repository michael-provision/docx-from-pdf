# -*- coding: utf-8 -*-

'''Page object parsed with PDF raw dict.

In addition to base structure described in :py:class:`~docx_from_pdf.page.RawPage`,
some new features, e.g. sections, table block, are also included.
Page elements structure:

* :py:class:`~docx_from_pdf.page.Page` >> :py:class:`~docx_from_pdf.layout.Section` >> :py:class:`~docx_from_pdf.layout.Column`
    * :py:class:`~docx_from_pdf.layout.Blocks`
        * :py:class:`~docx_from_pdf.text.TextBlock` >>
          :py:class:`~docx_from_pdf.text.Line` >>
          :py:class:`~docx_from_pdf.text.TextSpan` / :py:class:`~docx_from_pdf.image.ImageSpan` >>
          :py:class:`~docx_from_pdf.text.Char`
        * :py:class:`~docx_from_pdf.table.TableBlock` >>
          :py:class:`~docx_from_pdf.table.Row` >>
          :py:class:`~docx_from_pdf.table.Cell`
            * :py:class:`~docx_from_pdf.layout.Blocks`
            * :py:class:`~docx_from_pdf.shape.Shapes`
    * :py:class:`~docx_from_pdf.shape.Shapes`
        * :py:class:`~docx_from_pdf.shape.Shape.Stroke`
        * :py:class:`~docx_from_pdf.shape.Shape.Fill`
        * :py:class:`~docx_from_pdf.shape.Shape.Hyperlink`

::

    {
        "id": 0, # page index
        "width" : w,
        "height": h,
        "margin": [left, right, top, bottom],
        "sections": [{
            ... # section properties
        }, ...],
        "floats": [{
            ... # floating picture
        }, ...]
    }

'''

from docx.shared import Pt
from docx.enum.section import WD_SECTION
from ..common.Collection import BaseCollection
from ..common.share import debug_plot
from .BasePage import BasePage
from ..layout.Sections import Sections
from ..image.ImageBlock import ImageBlock
from ..common import constants
from ..common.docx import reset_paragraph_format
from .Furniture import (FurnitureContainer, FurnitureLayout, footer_settle,
    insert_page_number_fields, remove_trailing_spacers)


class Page(BasePage):
    '''Object representing the whole page, e.g. margins, sections.'''

    def __init__(self, id:int=-1,
                        skip_parsing:bool=True,
                        width:float=0.0,
                        height:float=0.0,
                        header=None,
                        footer=None,
                        margin:tuple=None,
                        sections:Sections=None,
                        float_images:BaseCollection=None):
        '''Initialize page layout.

        Args:
            id (int, optional): Page index. Defaults to -1.
            skip_parsing (bool, optional): Don't parse page if True. Defaults to True.
            width (float, optional): Page width. Defaults to 0.0.
            height (float, optional): Page height. Defaults to 0.0.
            header (FurnitureLayout, optional): Page header layout. Defaults to None.
            footer (FurnitureLayout, optional): Page footer layout. Defaults to None.
            margin (tuple, optional): Page margin. Defaults to None.
            sections (Sections, optional): Page contents. Defaults to None.
            float_images (BaseCollection, optional): Float images in th is page. Defaults to None.
        '''
        # page index
        self.id = id
        self.skip_parsing = skip_parsing

        # page size and margin
        super().__init__(width=width, height=height, margin=margin)

        # flow structure:
        # Section -> Column -> Blocks -> TextBlock/TableBlock
        # TableBlock -> Row -> Cell -> Blocks
        self.sections = sections or Sections(parent=self)

        # page header, footer
        self.header = header
        self.footer = footer
        self.page_number_fields = False

        # floating images are separate node under page
        self.float_images = float_images or BaseCollection()

        self._finalized = False


    @property
    def finalized(self): return self._finalized


    def store(self):
        '''Store parsed layout in dict format.'''
        res = {
            'id'      : self.id,
            'width'   : self.width,
            'height'  : self.height,
            'margin'  : self.margin,
            'sections': self.sections.store(),
            'header'  : self.header.store() if self.header else None,
            'footer'  : self.footer.store() if self.footer else None,
            'page_number_fields': self.page_number_fields,
            'floats'  : self.float_images.store()
        }
        return res


    def restore(self, data:dict):
        '''Restore Layout from parsed results.'''
        # page id
        self.id = data.get('id', -1)

        # page width/height
        self.width = data.get('width', 0.0)
        self.height = data.get('height', 0.0)
        self.margin = data.get('margin', (0,) * 4)

        # parsed layout
        self.sections.restore(data.get('sections', []))
        self.header = FurnitureLayout.restore(data['header']) if data.get('header') else None
        self.footer = FurnitureLayout.restore(data['footer']) if data.get('footer') else None
        self.page_number_fields = data.get('page_number_fields', False)

        # float images
        self._restore_float_images(data.get('floats', []))

        # Suppose layout is finalized when restored; otherwise, set False explicitly
        # out of this method.
        self._finalized = True

        return self


    @debug_plot('Final Layout')
    def parse(self, **settings):
        '''Parse page layout.'''
        self.sections.parse(**settings)
        for layout in (self.header, self.footer):
            if layout: layout.parse(**settings)
        self._finalized = True
        return self.sections # for debug plot


    def extract_tables(self, **settings):
        '''Extract content from tables (top layout only).

        .. note::
            Before running this method, the page layout must be either parsed from source
            page or restored from parsed data.
        '''
        # table blocks
        collections = []
        for section in self.sections:
            for column in section:
                if settings['extract_stream_table']:
                    collections.extend(column.blocks.table_blocks)
                else:
                    collections.extend(column.blocks.lattice_table_blocks)

        # check table
        tables = [] # type: list[ list[list[str]] ]
        for table_block in collections:
            tables.append(table_block.text)

        return tables


    def make_docx(self, doc):
        '''Set page size, margin, and create page.

        .. note::
            Before running this method, the page layout must be either parsed from source
            page or restored from parsed data.

        Args:
            doc (Document): ``python-docx`` document object
        '''
        # new page
        # count pages rather than paragraphs: a page whose content all moved to its
        # header and footer adds no body paragraph
        pages_made = getattr(doc, '_docx_from_pdf_pages', 0)
        doc._docx_from_pdf_pages = pages_made + 1
        if pages_made:
            section = doc.add_section(WD_SECTION.NEW_PAGE)
            # the paragraph holding the previous page's section break must not need room
            # of its own, or a page filled to its footer spills onto an extra blank page
            reset_paragraph_format(doc.paragraphs[-1], Pt(constants.MIN_LINE_SPACING))
            remove_trailing_spacers(doc.paragraphs[-1]._p)
        else:
            section = doc.sections[0] # a default section is there when opening docx

        # page size
        section.page_width  = Pt(self.width)
        section.page_height = Pt(self.height)

        # page margin
        left,right,top,bottom = self.margin
        section.left_margin = Pt(left)
        section.right_margin = Pt(right)
        section.top_margin = Pt(top)
        section.bottom_margin = Pt(bottom)

        # header and footer
        self._make_furniture(doc, section, 'header')
        self._make_furniture(doc, section, 'footer')

        # create flow layout: sections
        self.sections.make_docx(doc)


    def _make_furniture(self, doc, section, name:str):
        '''Write the page header or footer into the Word section that holds this page.'''
        layout = getattr(self, name)
        part = getattr(section, name)
        if not layout:
            # a section inherits the previous one's part unless it has its own
            if any(not getattr(earlier, name).is_linked_to_previous for earlier in doc.sections[:-1]):
                part.is_linked_to_previous = False
            return
        x0, y0, x1, y1 = layout.bbox
        if name == 'header':
            section.header_distance = Pt(y0)
        else:
            section.footer_distance = Pt(max(self.height - y1 - footer_settle(self.height, y1), 0))
        part.is_linked_to_previous = False
        placeholder = part.paragraphs[0]._p if part.paragraphs else None
        layout.make_docx(FurnitureContainer(part, Pt(x1 - x0)))
        if placeholder is not None and len(part._element) > 1 and not placeholder.xpath('.//w:t'):
            part._element.remove(placeholder)
        if self.page_number_fields:
            insert_page_number_fields(part._element)


    def _restore_float_images(self, raws:list):
        '''Restore float images.'''
        self.float_images.reset()
        for raw in raws:
            image = ImageBlock(raw)
            image.set_float_image_block()
            self.float_images.append(image)
