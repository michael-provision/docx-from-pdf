# -*- coding: utf-8 -*-

'''Row in a table.
'''

from copy import deepcopy

from docx.enum.table import WD_ROW_HEIGHT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from .Cells import Cells
from ..common.Element import Element
from ..common import docx


class Row(Element):
    '''Row in a table.'''
    def __init__(self, raw:dict=None):
        if raw is None: raw = {}
        super().__init__(raw)

        # logical row height
        self.height = raw.get('height', 0.0)

        # cells in row
        self._cells = Cells(parent=self).restore(raw.get('cells', []))


    def __getitem__(self, idx):
        try:
            cell = self._cells[idx]
        except IndexError:
            msg = f'Cell index {idx} out of range'
            raise IndexError(msg)
        else:
            return cell

    def __iter__(self):
        return (cell for cell in self._cells)

    def __len__(self):
        return len(self._cells)


    def append(self, cell):
        '''Append cell to row and update bbox accordingly.'''
        self._cells.append(cell)


    def store(self):
        res = super().store()
        res.update({
            'height': self.height,
            'cells': self._cells.store()
        })

        return res


    def make_docx(self, table, idx_row:int):
        '''Create row of docx table.
        
        Args:
            table (Table): ``python-docx`` table instance.
            idx_row (int): Current row index.
        '''  
        # set row height
        docx_row = table.rows[idx_row]

        # to control the layout precisely, set `exact` value, rather than `at least` value
        # the associated steps in MS word: Table Properties -> Row -> Row height -> exactly
        docx_row.height_rule = WD_ROW_HEIGHT.EXACTLY

        if self.height < 0: # to prevent negative height validation
            self.height = 0

        # NOTE: row height is counted from center-line of top border to center line of bottom border
        docx_row.height = Pt(self.height)

        # set cell style and contents
        for idx_col in range(len(table.columns)):
            self._cells[idx_col].make_docx(table, (idx_row, idx_col))

        Row._match_empty_cells_to_row_text(docx_row)


    @staticmethod
    def _match_empty_cells_to_row_text(docx_row):
        '''Give empty cells the paragraph spacing and run size of the row's text.

        Otherwise the empty paragraph Word requires in each cell takes the
        document's default spacing, which sets the row height once the height is
        no longer exact, and text typed into the cell (e.g. an INCLUDE mark)
        would not match the row.
        '''
        cells = list({id(cell._tc): cell for cell in docx_row.cells}.values())
        sample = next((paragraph for cell in cells for paragraph in cell.paragraphs
                       if paragraph.text.strip()), None)
        if sample is None: return
        sample_spacing = sample._p.pPr.find(qn('w:spacing')) if sample._p.pPr is not None else None
        sample_run = next((run for run in sample.runs if run.text.strip()), None)
        sample_rpr = sample_run._r.rPr if sample_run is not None else None
        for cell in cells:
            # only the bare paragraph Word requires: images or nested tables keep their own layout
            if len(cell._tc.findall(qn('w:p'))) != 1 or cell._tc.find(qn('w:tbl')) is not None \
                or cell.paragraphs[0]._p.find(qn('w:r')) is not None: continue
            pPr = cell.paragraphs[0]._p.get_or_add_pPr()
            if sample_spacing is not None:
                docx.set_child_in_order(pPr, deepcopy(sample_spacing), docx.PARAGRAPH_PROPERTY_ORDER)
            if sample_rpr is not None:
                mark = OxmlElement('w:rPr')
                for tag in ('w:rFonts', 'w:sz', 'w:szCs'):
                    element = sample_rpr.find(qn(tag))
                    if element is not None:
                        docx.set_child_in_order(mark, deepcopy(element), docx.RUN_PROPERTY_ORDER)
                if len(mark):
                    docx.set_child_in_order(pPr, mark, docx.PARAGRAPH_PROPERTY_ORDER)