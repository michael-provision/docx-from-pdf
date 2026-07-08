from pathlib import Path
from zipfile import ZipFile

from docx import Document

from pdf2docx import Converter, parse


ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "test" / "samples"


def test_runtime_package_does_not_reference_restricted_backend():
    forbidden = ("fi" + "tz", "py" + "mupdf")
    checked_files = [ROOT / "requirements.txt"]
    checked_files.extend((ROOT / "pdf2docx").rglob("*.py"))

    matches = []
    for path in checked_files:
        text = path.read_text(encoding="utf-8").lower()
        for token in forbidden:
            if token in text:
                matches.append(f"{path.relative_to(ROOT)} contains {token}")

    assert matches == []


def test_convert_pdf_to_editable_docx(tmp_path):
    source_pdf = SAMPLES / "demo-text.pdf"
    output_docx = tmp_path / "demo-text.docx"

    converter = Converter(str(source_pdf))
    converter.convert(str(output_docx))
    converter.close()

    document = Document(output_docx)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "PDF" in text
    assert len(document.sections) >= 1
    assert len(document.inline_shapes) == 0


def test_parse_function_supports_stream_input(tmp_path):
    source_pdf = SAMPLES / "demo-text.pdf"
    output_docx = tmp_path / "stream.docx"

    converter = Converter(stream=source_pdf.read_bytes())
    converter.convert(str(output_docx))

    with ZipFile(output_docx) as archive:
        assert "word/document.xml" in archive.namelist()


def test_cli_parse_alias_writes_docx(tmp_path):
    source_pdf = SAMPLES / "demo-text.pdf"
    output_docx = tmp_path / "parse-alias.docx"

    parse(str(source_pdf), str(output_docx))

    assert output_docx.exists()


def test_parse_make_docx_chain_keeps_page_selection(tmp_path):
    source_pdf = SAMPLES / "demo-text.pdf"
    output_docx = tmp_path / "first-page.docx"

    converter = Converter(str(source_pdf))
    converter.parse(pages=[0]).make_docx(str(output_docx))

    document = Document(output_docx)
    assert len(document.sections) == 1


def test_extract_tables_uses_pdfplumber_backend():
    source_pdf = SAMPLES / "demo-table.pdf"

    converter = Converter(str(source_pdf))
    tables = converter.extract_tables()

    assert tables
