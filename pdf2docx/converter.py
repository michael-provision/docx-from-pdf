"""PDF to DOCX converter using pdfplumber/pdfminer.

This fork keeps the public ``Converter`` API but replaces the previous
PDF-specific layout engine with pdfplumber/pdfminer extraction and python-docx
generation. It is intentionally conservative: the output is editable text laid
out with page sections, paragraph offsets, indentation, tab stops, and basic
font styling. It does not attempt to reconstruct vector drawings or image-only
pages yet.
"""

from __future__ import annotations

import json
import logging
import os
from io import BytesIO
from time import perf_counter
from typing import AnyStr, IO, Iterable, Union

import pdfplumber
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Pt

logging.basicConfig(level=logging.INFO, format="[%(levelname)s] %(message)s")

DEFAULT_LINE_TOLERANCE = 3.0
DEFAULT_GAP_TOLERANCE = 8.0


class Converter:
    """Convert PDF pages to DOCX using permissively licensed PDF tooling."""

    def __init__(
        self,
        pdf_file: str | None = None,
        password: str | None = None,
        stream: bytes | None = None,
    ):
        if not pdf_file and not stream:
            raise ValueError("Either pdf_file or stream must be given.")
        self.filename_pdf = pdf_file
        self.password = str(password or "")
        self._stream = stream
        self._last_result: dict = {}
        self._last_page_indexes: list[int] | None = None

    @property
    def default_settings(self) -> dict:
        return {
            "line_tolerance": DEFAULT_LINE_TOLERANCE,
            "gap_tolerance": DEFAULT_GAP_TOLERANCE,
            "preserve_page_size": True,
            "raw_exceptions": False,
            "multi_processing": False,
            "zero_based_index": True,
        }

    @property
    def pages(self) -> list[dict]:
        return self._last_result.get("pages", [])

    def close(self) -> None:
        """Retained for API compatibility; pdfplumber files are opened per call."""

    def convert(
        self,
        docx_filename: Union[str, IO[AnyStr], None] = None,
        start: int = 0,
        end: int | None = None,
        pages: list | None = None,
        **kwargs,
    ) -> None:
        """Convert specified PDF pages to a DOCX file."""
        start_time = perf_counter()
        logging.info("Start to convert %s", self.filename_pdf or "<stream>")
        settings = self.default_settings
        settings.update(kwargs)
        if settings["multi_processing"]:
            raise ConversionException("multi_processing is not supported by this backend.")

        if docx_filename is None:
            docx_filename = self._default_docx_filename()

        with self._open_pdf() as pdf:
            page_indexes = self._page_indexes(start, end, pages, len(pdf.pages))
            self._last_page_indexes = page_indexes
            document = Document()
            summaries = []
            for output_page_number, page_index in enumerate(page_indexes):
                page = pdf.pages[page_index]
                logging.info("(%d/%d) Page %d", output_page_number + 1, len(page_indexes), page_index + 1)
                try:
                    summary = self._append_pdf_page(document, page, output_page_number, settings)
                except Exception as exc:
                    if settings["raw_exceptions"]:
                        raise
                    logging.error("Ignore page %d due to conversion error: %s", page_index + 1, exc)
                    continue
                summaries.append(summary | {"id": page_index})

            document.save(docx_filename)
            self._last_result = {
                "filename": os.path.basename(self.filename_pdf or "stream.pdf"),
                "page_cnt": len(pdf.pages),
                "pages": summaries,
            }
        logging.info("Terminated in %.2fs.", perf_counter() - start_time)

    def parse(self, start: int = 0, end: int | None = None, pages: list | None = None, **kwargs):
        """Collect page extraction summaries for compatibility with old callers."""
        settings = self.default_settings
        settings.update(kwargs)
        with self._open_pdf() as pdf:
            page_indexes = self._page_indexes(start, end, pages, len(pdf.pages))
            self._last_page_indexes = page_indexes
            summaries = []
            for page_index in page_indexes:
                page = pdf.pages[page_index]
                words = self._extract_words(page)
                summaries.append(
                    {
                        "id": page_index,
                        "width": page.width,
                        "height": page.height,
                        "word_count": len(words),
                        "line_count": len(self._cluster_lines(words, settings["line_tolerance"])),
                    }
                )
            self._last_result = {
                "filename": os.path.basename(self.filename_pdf or "stream.pdf"),
                "page_cnt": len(pdf.pages),
                "pages": summaries,
            }
        return self

    def make_docx(self, filename_or_stream=None, **kwargs) -> None:
        """Create a DOCX from the original source; parse-tree replay is not supported."""
        if self._last_page_indexes is not None and "pages" not in kwargs:
            kwargs["pages"] = self._last_page_indexes
        self.convert(filename_or_stream, **kwargs)

    def extract_tables(
        self,
        start: int = 0,
        end: int | None = None,
        pages: list | None = None,
        **kwargs,
    ) -> list:
        settings = self.default_settings
        settings.update(kwargs)
        tables = []
        with self._open_pdf() as pdf:
            page_indexes = self._page_indexes(start, end, pages, len(pdf.pages))
            for page_index in page_indexes:
                tables.extend(pdf.pages[page_index].extract_tables() or [])
        return tables

    def debug_page(
        self,
        i: int,
        docx_filename: str | None = None,
        debug_pdf: str | None = None,
        layout_file: str | None = None,
        **kwargs,
    ) -> None:
        """Convert one page and serialize the lightweight extraction summary."""
        self.convert(docx_filename, pages=[i], **kwargs)
        if layout_file:
            self.serialize(layout_file)

    def store(self) -> dict:
        return self._last_result

    def restore(self, data: dict):
        self._last_result = data
        return self

    def serialize(self, filename: str) -> None:
        with open(filename, "w", encoding="utf-8") as file:
            json.dump(self.store(), file, indent=4)

    def deserialize(self, filename: str):
        with open(filename, "r", encoding="utf-8") as file:
            self.restore(json.load(file))
        return self

    def _open_pdf(self):
        if self._stream is not None:
            return pdfplumber.open(BytesIO(self._stream), password=self.password or None)
        return pdfplumber.open(self.filename_pdf, password=self.password or None)

    def _default_docx_filename(self) -> str:
        if not self.filename_pdf:
            raise ConversionException("Please specify a docx file name for stream input.")
        return f"{self.filename_pdf[0:-len('.pdf')]}.docx"

    def _append_pdf_page(self, document: Document, page, output_page_number: int, settings: dict) -> dict:
        self._configure_section(document, output_page_number, page, settings)
        words = self._extract_words(page)
        lines = self._cluster_lines(words, settings["line_tolerance"])
        previous_bottom = 0.0
        for line in lines:
            paragraph = document.add_paragraph()
            self._format_paragraph(paragraph, line, previous_bottom, settings)
            self._append_line_runs(paragraph, line, settings)
            previous_bottom = line["bottom"]
        return {
            "width": page.width,
            "height": page.height,
            "word_count": len(words),
            "line_count": len(lines),
        }

    @staticmethod
    def _configure_section(document: Document, output_page_number: int, page, settings: dict) -> None:
        if output_page_number == 0:
            section = document.sections[0]
        else:
            section = document.add_section(WD_SECTION.NEW_PAGE)
        if settings["preserve_page_size"]:
            section.page_width = Pt(page.width)
            section.page_height = Pt(page.height)
        section.top_margin = Pt(0)
        section.bottom_margin = Pt(0)
        section.left_margin = Pt(0)
        section.right_margin = Pt(0)
        section.header_distance = Pt(0)
        section.footer_distance = Pt(0)

    @staticmethod
    def _extract_words(page) -> list[dict]:
        return page.extract_words(
            extra_attrs=["fontname", "size"],
            keep_blank_chars=False,
            use_text_flow=False,
        )

    @staticmethod
    def _cluster_lines(words: Iterable[dict], line_tolerance: float) -> list[dict]:
        lines: list[dict] = []
        for word in sorted(words, key=lambda item: (item["top"], item["x0"])):
            for line in reversed(lines[-6:]):
                if abs(line["top"] - word["top"]) <= line_tolerance:
                    line["words"].append(word)
                    line["top"] = min(line["top"], word["top"])
                    line["bottom"] = max(line["bottom"], word["bottom"])
                    break
            else:
                lines.append({"top": word["top"], "bottom": word["bottom"], "words": [word]})
        for line in lines:
            line["words"].sort(key=lambda item: item["x0"])
        return sorted(lines, key=lambda item: (item["top"], item["words"][0]["x0"]))

    @staticmethod
    def _format_paragraph(paragraph, line: dict, previous_bottom: float, settings: dict) -> None:
        first_word = line["words"][0]
        paragraph_format = paragraph.paragraph_format
        paragraph_format.space_before = Pt(max(0.0, line["top"] - previous_bottom))
        paragraph_format.space_after = Pt(0)
        paragraph_format.left_indent = Pt(max(0.0, first_word["x0"]))
        paragraph_format.line_spacing = 1

    @staticmethod
    def _append_line_runs(paragraph, line: dict, settings: dict) -> None:
        words = line["words"]
        line_left = words[0]["x0"]
        previous_x1 = line_left
        for word_index, word in enumerate(words):
            if word_index:
                gap = word["x0"] - previous_x1
                if gap > settings["gap_tolerance"]:
                    paragraph.paragraph_format.tab_stops.add_tab_stop(
                        Pt(max(0.0, word["x0"] - line_left))
                    )
                    paragraph.add_run("\t")
                else:
                    paragraph.add_run(" ")
            run = paragraph.add_run(word["text"])
            Converter._style_run(run, word)
            previous_x1 = word["x1"]

    @staticmethod
    def _style_run(run, word: dict) -> None:
        size = word.get("size")
        if size:
            run.font.size = Pt(float(size))
        font_name = word.get("fontname") or ""
        run.bold = "Bold" in font_name or "Black" in font_name
        run.italic = "Italic" in font_name or "Oblique" in font_name

    @staticmethod
    def _page_indexes(start, end, pages, pdf_len: int):
        if pages:
            return [int(index) for index in pages]
        actual_end = end if end is not None else pdf_len
        return list(range(pdf_len)[slice(int(start), int(actual_end))])

    @staticmethod
    def _color_output(msg: str) -> str:
        return f"\033[1;36m{msg}\033[0m"


class ConversionException(Exception):
    pass


class MakedocxException(ConversionException):
    pass
